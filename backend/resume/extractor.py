"""
Resume text extraction.

Takes an uploaded PDF or DOCX file and pulls out plain text so it can be
handed to Gemini for scoring. No AI here - just parsing.
"""

import io

from pypdf import PdfReader
from docx import Document


class UnsupportedFileTypeError(ValueError):
    pass


def extract_text(file_stream, filename):
    """
    file_stream: a file-like object (e.g. Flask's request.files['resume'].stream)
    filename: original filename, used only to decide PDF vs DOCX
    """
    lower = filename.lower()

    if lower.endswith(".pdf"):
        return _extract_pdf(file_stream)
    elif lower.endswith(".docx"):
        return _extract_docx(file_stream)
    else:
        raise UnsupportedFileTypeError(
            f"Unsupported file type: {filename}. Please upload a .pdf or .docx file."
        )


def _extract_pdf(file_stream):
    reader = PdfReader(file_stream)
    pages_text = []
    for page in reader.pages:
        text = page.extract_text() or ""
        pages_text.append(text)
    return "\n".join(pages_text).strip()


def _extract_docx(file_stream):
    # python-docx needs a file-like object with .seek support; Flask gives us that.
    buffer = io.BytesIO(file_stream.read())
    document = Document(buffer)
    paragraphs = [p.text for p in document.paragraphs]

    # Also pull text out of any tables (some resumes use table layouts)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    paragraphs.append(cell.text)

    return "\n".join(paragraphs).strip()
